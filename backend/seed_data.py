import sys
from pathlib import Path
import openpyxl

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Support running directly or as a module
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"

def seed_events():
    """Seeds the events table from data/events_seed.xlsx if the table is currently empty."""
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) AS count FROM events;")
    row_count = cursor.fetchone()["count"]

    if row_count > 0:
        print(f"Database already contains {row_count} events. Skipping seeding.")
        conn.close()
        return

    if not SEED_FILE_PATH.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed file not found at '{SEED_FILE_PATH}'. "
            "Please copy the companion data file 'events_seed.xlsx' into the data/ folder before running this script."
        )

    try:
        workbook = openpyxl.load_workbook(SEED_FILE_PATH, data_only=True)
    except Exception as exc:
        conn.close()
        raise RuntimeError(f"Failed to open '{SEED_FILE_PATH}': {exc}") from exc

    if "Events" not in workbook.sheetnames:
        conn.close()
        raise ValueError(f"Sheet 'Events' not found in '{SEED_FILE_PATH}'. Found sheets: {workbook.sheetnames}")

    sheet = workbook["Events"]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        conn.close()
        raise ValueError(f"Seed file '{SEED_FILE_PATH}' contains no data rows.")

    header = [str(col).strip() if col is not None else "" for col in rows[0]]
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]

    if [h.lower() for h in header] != [h.lower() for h in expected_header]:
        conn.close()
        raise ValueError(
            f"Header mismatch in '{SEED_FILE_PATH}'.\nExpected: {expected_header}\nGot: {header}"
        )

    events_inserted = 0
    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or all(v is None for v in row):
            continue

        event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            int(event_id) if event_id is not None else None,
            str(title) if title is not None else "",
            str(category) if category is not None else "",
            str(description) if description is not None else "",
            str(event_date) if event_date is not None else "",
            str(event_time) if event_time is not None else "",
            str(venue) if venue is not None else "",
            int(capacity) if capacity is not None else 0,
            int(fee) if fee is not None else 0,
            str(icon) if icon is not None else ""
        ))
        events_inserted += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {events_inserted} events from '{SEED_FILE_PATH}'.")

if __name__ == "__main__":
    seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC;")
    rows = cursor.fetchall()
    print(f"\n--- Events in Database ({len(rows)} total) ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} | {r['category']} | {r['event_date']} {r['event_time']} | {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
